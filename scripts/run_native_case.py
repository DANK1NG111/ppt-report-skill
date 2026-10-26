"""Retained TEST FIXTURE demonstrates separately editable native evidence."""

import json
from pathlib import Path
import sys
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from ppt_report_skill.pipeline import create_presentation
from pptx import Presentation

out = Path("output/native_case")
out.mkdir(parents=True, exist_ok=True)
with ZipFile("assets/templates/base_template.pptx") as z:
    image = out / "reference.jpeg"
    image.write_bytes(z.read("ppt/media/image29.jpeg"))
request = {
    "topic": "TEST FIXTURE 原生对象验证",
    "presentation_type": "custom",
    "required_sections": ["阶段数据", "执行流程", "现场图片"],
    "visuals": [
        {
            "type": "table",
            "section": "阶段数据",
            "title": "测试结果",
            "data": {"rows": [["项目", "数量"], ["通过", "16"], ["待处理", "2"]]},
        },
        {
            "type": "chart",
            "section": "阶段数据",
            "title": "监测点数量",
            "data": {
                "kind": "column",
                "title": "监测点数量",
                "unit": "个",
                "time": "2026年9月",
                "categories": ["甲", "乙", "丙"],
                "series": [{"name": "数量", "values": [12, 18, 9]}],
            },
        },
        {
            "type": "diagram",
            "section": "执行流程",
            "title": "执行流程",
            "data": {
                "kind": "process",
                "nodes": [
                    {"id": "a", "label": "收集材料"},
                    {"id": "b", "label": "核对事实"},
                    {"id": "c", "label": "交付成果"},
                ],
            },
        },
        {
            "type": "image",
            "section": "现场图片",
            "title": "材料原图",
            "path": str(image.resolve()),
        },
    ],
}
result = create_presentation(
    "assets/templates/base_template.pptx",
    [],
    request,
    out / "final_presentation.pptx",
    direct_text="# 阶段数据\nTEST FIXTURE：2026年9月共完成18项测试，其中16项通过、2项待处理。\n# 执行流程\n执行流程包括收集材料、核对事实、交付成果三个阶段。\n# 现场图片\n图片来自模板素材，保留为独立对象。",
)
p = Presentation(result["path"])
assert any(s.has_table for sl in p.slides for s in sl.shapes)
assert any(s.has_chart for sl in p.slides for s in sl.shapes)
assert sum(s._element.tag.endswith("cxnSp") for sl in p.slides for s in sl.shapes) >= 2
print(
    json.dumps(
        {k: result[k] for k in ["path", "slide_count", "build", "status"]},
        ensure_ascii=False,
    )
)
