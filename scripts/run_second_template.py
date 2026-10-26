"""Create a separate 4:3 TEST FIXTURE template and exercise the full pipeline."""

from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from ppt_report_skill.pipeline import create_presentation
from ppt_report_skill.models import save_json


def make_template(path):
    p = Presentation()
    p.slide_width = Inches(10)
    p.slide_height = Inches(7.5)

    def text(slide, value, x, y, w, h, size, color="17365D"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text = value
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(size)
                run.font.color.rgb = RGBColor.from_string(color)
        return shape

    for kind in ["cover", "contents", "content", "closing"]:
        s = p.slides.add_slide(p.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("F2F6FA")
        if kind in {"cover", "closing"}:
            text(
                s,
                "TEST FIXTURE 第二模板" if kind == "cover" else "感谢聆听",
                0.65,
                2.5,
                8.7,
                1.3,
                44,
            )
        else:
            text(
                s,
                "目录" if kind == "contents" else "正文标题",
                0.65,
                0.45,
                8.7,
                0.8,
                30,
            )
            if kind == "contents":
                for i, value in enumerate(["背景", "阶段成果", "下一步"]):
                    text(s, value, 1.0, 2 + i * 1.1, 7.8, 0.7, 24)
            else:
                text(
                    s,
                    "正文示例材料仅用于模板容量分析。\n保留原生文字并写入实际材料。",
                    0.65,
                    1.65,
                    8.7,
                    5.0,
                    22,
                    "202020",
                )
        text(s, "Copyright TEST FIXTURE", 0.65, 7.05, 8.5, 0.25, 12)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    p.save(path)
    return path


def run_case(output="output/second_template"):
    out = Path(output)
    template = make_template(out / "second_template.pptx")
    request = {
        "topic": "TEST FIXTURE 第二模板验证",
        "presentation_type": "custom",
        "required_sections": ["背景", "阶段成果", "下一步"],
        "target_slide_count": 6,
        "duration": 4,
    }
    material = "# 背景\nTEST FIXTURE：本案例用于验证第二模板，与真实业务无关。\n# 阶段成果\n截至2026年9月30日，完成率87.6%，接入12个监测点。\n# 下一步\n计划修复2项未通过测试，保留原始日志。"
    result = create_presentation(
        template, [], request, out / "final_presentation.pptx", direct_text=material
    )
    assert result["qa"]["errors"] == 0 and result["slide_count"] == 6
    assert Presentation(result["path"]).slide_width == Inches(10)
    save_json(
        {"status": "PASS", "template": str(template.resolve()), "result": result},
        out / "generalization_report.json",
    )
    print(
        "PASS: second template, 4:3 blue style, fixed order, exact facts, 6 pages rendered"
    )
    return result


if __name__ == "__main__":
    run_case()
