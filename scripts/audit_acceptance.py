"""Inspect final artifacts against source facts, plan, theme and native object contracts."""

from pathlib import Path
import json
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from ppt_report_skill.validator.content import slide_text, compact
from ppt_report_skill.template.analyzer import walk_shapes
from ppt_report_skill.models import save_json


def audit(receipt):
    r = json.loads(Path(receipt).read_text())
    plans = json.loads(Path(r["build"], "slide_plans.json").read_text())
    profile = json.loads(Path(r["build"], "template_profile.json").read_text())
    facts = json.loads(Path(r["build"], "source_facts.json").read_text())
    request = json.loads(Path(r["build"], "request.json").read_text())
    prs = Presentation(r["path"])
    original = Presentation(r["template"])
    text = compact("\n".join(slide_text(s) for s in prs.slides))
    layouts = {l["layout_id"]: l for l in profile["layouts"]}
    assert all(compact(f["content"]) in text for f in facts), "source fact missing"
    rendered_text = compact("\n".join(r["rendered"]["page_texts"]))
    assert all(compact(f["content"]) in rendered_text for f in facts), (
        "rendered source fact missing"
    )
    order = []
    for p in plans:
        if p["slide_type"] in {"content", "section"} and (
            not order or order[-1] != p["section"]
        ):
            order.append(p["section"])
    if request["required_sections"]:
        assert order == request["required_sections"], "section order"
    with ZipFile(r["template"]) as a, ZipFile(r["path"]) as b:
        names = [
            n
            for n in a.namelist()
            if n.endswith(".xml")
            and n.startswith(("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/"))
        ]
        assert all(
            etree.tostring(etree.fromstring(a.read(n)), method="c14n")
            == etree.tostring(etree.fromstring(b.read(n)), method="c14n")
            for n in names
        ), "theme/master/layout changed"
    fixed = 0
    for slide, plan in zip(prs.slides, plans):
        l = layouts[plan["layout_id"]]
        old = {
            s.shape_id: s
            for s in walk_shapes(original.slides[l["source_slide"] - 1].shapes)
        }
        new = {s.shape_id: s for s in walk_shapes(slide.shapes)}
        for b in l["fixed_shapes"]:
            source = old[b["shape_id"]]
            target = new[b["shape_id"]]
            assert (source.left, source.top, source.width, source.height) == (
                target.left,
                target.top,
                target.width,
                target.height,
            )
            if source.has_text_frame:
                assert source.text == target.text
            if hasattr(source, "image"):
                assert source.image.sha1 == target.image.sha1
            fixed += 1
    objects = {
        "text": 0,
        "tables": 0,
        "charts": 0,
        "pictures": 0,
        "connectors": 0,
        "shapes": 0,
    }
    for slide in prs.slides:
        for s in walk_shapes(slide.shapes):
            objects["text"] += int(s.has_text_frame)
            objects["tables"] += int(s.has_table)
            objects["charts"] += int(s.has_chart)
            objects["pictures"] += int(hasattr(s, "image"))
            objects["connectors"] += int(s._element.tag.endswith("cxnSp"))
            objects["shapes"] += int(s._element.tag.endswith("sp"))
    return {
        "status": "PASS",
        "sha256": r["sha256"],
        "path": r["path"],
        "slide_count": len(prs.slides),
        "source_facts": len(facts),
        "source_facts_exact": True,
        "section_sequence": order,
        "theme_master_layout_preserved": True,
        "fixed_objects_checked": fixed,
        "editable_objects": objects,
        "errors": r["qa"]["errors"],
        "warnings": r["qa"]["warnings"],
    }


if __name__ == "__main__":
    inputs = [
        *Path("output/real_template_cases").glob("*/pipeline_receipt.json"),
        Path("output/native_case/pipeline_receipt.json"),
        Path("output/second_template/pipeline_receipt.json"),
    ]
    reports = [audit(p) for p in inputs]
    assert all(r["errors"] == 0 for r in reports)
    save_json(
        {"status": "PASS", "cases": reports}, "output/acceptance/object_audit.json"
    )
    print(
        "PASS: exact facts, continuous section order, unchanged themes/masters/layouts and fixed objects; native evidence inspected"
    )
