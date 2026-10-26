"""Ordered, non-executing loaders. Source locations survive normalization."""

from pathlib import Path
import hashlib
import re
import fitz
from ppt_report_skill.models import MaterialBlock, ReportError
from ppt_report_skill.template.analyzer import walk_shapes
from .normalizer import normalize


def text_blocks(text, source="<direct text>"):
    result = []
    section = ""
    for n, line in enumerate(normalize(text).splitlines(), 1):
        if not line.strip():
            continue
        heading = re.match(
            r"^#{1,6}\s+(.+)|^(?:[一二三四五六七八九十]+[、.]|\d+[.、])\s*(.{1,30})$",
            line,
        )
        if heading:
            section = next(x for x in heading.groups() if x)
            kind = "heading"
        else:
            kind = "list" if re.match(r"^[-*+]\s", line) else "paragraph"
        result.append(
            MaterialBlock(
                str(source),
                f"line {n}",
                kind,
                re.sub(r"^#{1,6}\s+|^[-*+]\s+", "", line),
                section,
            )
        )
    return result


def load_material(path, asset_dir=None):
    path = Path(path).resolve()
    if not path.is_file():
        raise ReportError("MATERIAL_NOT_FOUND", f"材料不存在：{path}")
    suffix = path.suffix.lower()
    source = str(path)
    blocks = []

    def add(location, kind, text, section="", rows=None):
        text = normalize(text)
        if text:
            blocks.append(
                MaterialBlock(source, location, kind, text, section, rows or [])
            )

    def image(blob, ext, location, section=""):
        if asset_dir:
            dest = Path(asset_dir) / f"{hashlib.sha256(blob).hexdigest()[:20]}.{ext}"
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(blob)
            blocks.append(
                MaterialBlock(
                    source,
                    location,
                    "image",
                    "",
                    section,
                    image_path=str(dest.resolve()),
                )
            )

    try:
        if suffix in {".txt", ".md"}:
            raw = path.read_bytes()
            for encoding in ["utf-8-sig", "utf-16", "gb18030"]:
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeError:
                    continue
            else:
                raise ReportError("UNSUPPORTED_ENCODING", "无法识别材料编码。")
            blocks = text_blocks(text, source)
        elif suffix == ".docx":
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table

            document = Document(path)
            section = ""
            for n, element in enumerate(document.element.body, 1):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    paragraph = Paragraph(element, document)
                    kind = (
                        "heading"
                        if paragraph.style.name.startswith(("Heading", "标题"))
                        else "paragraph"
                    )
                    if kind == "heading":
                        section = paragraph.text
                    add(f"paragraph {n}", kind, paragraph.text, section)
                    for blip in element.xpath(".//a:blip"):
                        rid = blip.get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                        )
                        part = document.part.related_parts.get(rid)
                        if part:
                            image(
                                part.blob,
                                part.partname.ext,
                                f"paragraph {n} image",
                                section,
                            )
                elif tag == "tbl":
                    rows = [
                        [c.text for c in row.cells]
                        for row in Table(element, document).rows
                    ]
                    add(
                        f"table {n}",
                        "table",
                        "\n".join(" | ".join(r) for r in rows),
                        section,
                        rows,
                    )
        elif suffix == ".pdf":
            with fitz.open(path) as document:
                for i, page in enumerate(document, 1):
                    for j, b in enumerate(page.get_text("blocks", sort=True), 1):
                        if b[6] == 0:
                            add(f"page {i} block {j}", "paragraph", b[4])
                    if asset_dir:
                        for j, info in enumerate(page.get_images(full=True), 1):
                            data = document.extract_image(info[0])
                            image(data["image"], data["ext"], f"page {i} image {j}")
        elif suffix == ".pptx":
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            for i, slide in enumerate(Presentation(path).slides, 1):
                section = slide.shapes.title.text if slide.shapes.title else ""
                for shape in walk_shapes(slide.shapes):
                    location = f"slide {i} shape {shape.shape_id}"
                    if shape.has_text_frame:
                        add(
                            location,
                            "heading" if shape == slide.shapes.title else "paragraph",
                            shape.text,
                            section,
                        )
                    if shape.has_table:
                        rows = [[c.text for c in r.cells] for r in shape.table.rows]
                        add(
                            location,
                            "table",
                            "\n".join(" | ".join(r) for r in rows),
                            section,
                            rows,
                        )
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image(shape.image.blob, shape.image.ext, location, section)
                    if shape.has_chart:
                        for series in shape.chart.series:
                            add(
                                location,
                                "table",
                                f"{series.name}: "
                                + ", ".join(str(v) for v in series.values),
                                section,
                            )
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=False, read_only=True)
            for sheet in workbook:
                rows = [
                    [str(c.value) if c.value is not None else "" for c in row]
                    for row in sheet.iter_rows()
                ]
                rows = [r for r in rows if any(r)]
                add(
                    f"sheet {sheet.title}",
                    "table",
                    "\n".join(" | ".join(r) for r in rows),
                    sheet.title,
                    rows,
                )
            workbook.close()
        else:
            raise ReportError("UNSUPPORTED_FILE", f"不支持的材料格式：{suffix}")
    except ReportError:
        raise
    except Exception as exc:
        raise ReportError(
            "MATERIAL_READ_FAILED", f"材料读取失败：{path.name}", str(exc)
        ) from exc
    if not any(b.text for b in blocks):
        raise ReportError(
            "MATERIAL_EMPTY", f"材料无可提取文字：{path.name}；扫描 PDF 需先 OCR。"
        )
    return blocks


def load_materials(paths=(), direct_text=None, asset_dir=None):
    blocks = []
    for path in paths:
        blocks.extend(load_material(path, asset_dir))
    if direct_text:
        blocks.extend(text_blocks(direct_text))
    if not blocks:
        raise ReportError("MATERIAL_EMPTY", "请提供汇报材料。")
    return blocks
