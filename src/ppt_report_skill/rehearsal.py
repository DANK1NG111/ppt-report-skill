"""Self-contained offline rehearsal companion; editable PPTX stays the deliverable."""

import base64
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from .models import ReportError, TemplateProfile, SlidePlan
from .design import verify_design_contract


def export_rehearsal(receipt, output=None):
    result = (
        json.loads(Path(receipt).read_text())
        if isinstance(receipt, (str, Path))
        else receipt
    )
    if (
        hashlib.sha256(Path(result["path"]).read_bytes()).hexdigest()
        != result["sha256"]
    ):
        raise ReportError("STALE_REHEARSAL", "PPTX已改变，请重新渲染。")
    rendered = result["rendered"]
    if rendered["sha256"] != result["sha256"]:
        raise ReportError("STALE_REHEARSAL", "渲染与当前PPTX不匹配。")
    contract = json.loads(Path(result["build"], "design_spec.json").read_text())
    build = Path(result["build"])
    if result.get("design_lock_sha256") != contract.get("lock_sha256"):
        raise ReportError("STALE_REHEARSAL", "设计锁定与生成回执不一致。")
    verify_design_contract(
        contract,
        TemplateProfile.from_dict(
            json.loads((build / "template_profile.json").read_text())
        ),
        [SlidePlan(**p) for p in json.loads((build / "slide_plans.json").read_text())],
    )
    prs = Presentation(result["path"])
    pages = []
    if len(prs.slides) != len(rendered["images"]) or len(prs.slides) != len(
        contract["pages"]
    ):
        raise ReportError("STALE_REHEARSAL", "页面与设计记录不一致。")
    for i, (slide, path, plan) in enumerate(
        zip(prs.slides, rendered["images"], contract["pages"]), 1
    ):
        path = Path(path)
        if hashlib.sha256(path.read_bytes()).hexdigest() != rendered[
            "artifact_hashes"
        ].get(str(path)):
            raise ReportError("STALE_REHEARSAL", "页图已改变，请重新渲染。")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        pages.append(
            {
                "id": plan["slide_id"],
                "title": plan["title"],
                "purpose": plan["purpose"],
                "seconds": plan["estimated_seconds"],
                "notes": notes,
                "image": "data:image/png;base64,"
                + base64.b64encode(path.read_bytes()).decode(),
            }
        )
    data = (
        json.dumps(
            {
                "sha256": result["sha256"],
                "template": contract["template_sha256"],
                "width": prs.slide_width,
                "height": prs.slide_height,
                "pages": pages,
            },
            ensure_ascii=False,
        )
        .replace("<", "\\u003c")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )
    html = """<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>汇报排练</title><style>
*{box-sizing:border-box}body{margin:0;background:#f4f6f8;color:#172334;font:16px system-ui,sans-serif}header,footer{padding:12px 22px;display:flex;gap:14px;align-items:center;flex-wrap:wrap;background:white}header h1{font-size:18px;margin:0 auto 0 0}main{display:grid;grid-template-columns:minmax(0,3fr) minmax(270px,1fr);gap:20px;padding:20px}.stage img{width:100%;object-fit:contain;background:white;border:1px solid #ddd}aside textarea{width:100%;height:36vh;font:16px/1.6 system-ui;padding:12px;border:1px solid #bcc7d5}aside img{max-width:100%;max-height:21vh;object-fit:contain}h2{font-size:19px;margin:8px 0}button{font:inherit;padding:7px 14px;cursor:pointer}small{color:#506078}#time{font-variant-numeric:tabular-nums}#notice{padding:0 22px 14px} @media(max-width:800px){main{grid-template-columns:1fr}aside textarea{height:25vh}}
</style><header><h1>汇报排练</h1><span id="progress"></span><span id="time">00:00</span><button id="timer">开始计时</button><button id="reset">重置计时</button><label><input type="checkbox" id="auto">按计划时长自动翻页</label></header><main><div class="stage"><img id="current" alt="当前页"><h2 id="title"></h2><p id="purpose"></p></div><aside><h2>本页备注</h2><textarea id="notes" aria-label="本页备注"></textarea><small id="saved"></small><h2>下一页</h2><img id="next" alt="下一页"><p id="nexttitle"></p></aside></main><footer><button id="home">首页</button><button id="prev">上一页</button><button id="nextbutton">下一页</button><button id="end">尾页</button><span id="planned"></span></footer><p id="notice">离线排练辅助文件。幻灯片由PPTX实际渲染；主要对象请在PPTX中编辑。备注只保存到本机浏览器。</p><script id="deck-data" type="application/json">DATA</script><script>
'use strict';const deck=JSON.parse(document.getElementById('deck-data').textContent);const $=id=>document.getElementById(id);let index=0,running=false,total=0,pageElapsed=0,last=performance.now();const key=p=>'pptreport:notes:v1:'+deck.template+':'+p.id;function show(n){index=Math.max(0,Math.min(deck.pages.length-1,n));pageElapsed=0;const p=deck.pages[index];$('current').src=p.image;$('title').textContent=p.title;$('purpose').textContent=p.purpose;$('progress').textContent=(index+1)+' / '+deck.pages.length;$('planned').textContent='本页计划 '+Math.round(p.seconds)+' 秒；全稿 '+Math.round(deck.pages.reduce((s,p)=>s+p.seconds,0))+' 秒';try{$('notes').value=localStorage.getItem(key(p))??p.notes;$('saved').textContent='备注按稳定页面ID保存，重排不会串页。'}catch(e){$('notes').value=p.notes;$('saved').textContent='此浏览器不允许本地保存，备注仅保留本次页面。'}$('notes').scrollTop=0;const next=deck.pages[index+1];$('next').hidden=!next;if(next)$('next').src=next.image;$('nexttitle').textContent=next?next.title:'演示结束';$('prev').disabled=index===0;$('nextbutton').disabled=index===deck.pages.length-1;}
$('notes').addEventListener('input',()=>{try{localStorage.setItem(key(deck.pages[index]),$('notes').value);$('saved').textContent='已保存到本机。'}catch(e){$('saved').textContent='保存不可用，请复制备注到本地文件。'}});$('home').onclick=()=>show(0);$('prev').onclick=()=>show(index-1);$('nextbutton').onclick=()=>show(index+1);$('end').onclick=()=>show(deck.pages.length-1);$('timer').onclick=()=>{running=!running;last=performance.now();$('timer').textContent=running?'暂停计时':'继续计时'};$('reset').onclick=()=>{running=false;total=pageElapsed=0;$('timer').textContent='开始计时';$('time').textContent='00:00'};document.addEventListener('visibilitychange',()=>{last=performance.now()});document.addEventListener('keydown',e=>{if(e.target.tagName==='TEXTAREA'||e.target.tagName==='INPUT')return;if(e.key==='ArrowRight')show(index+1);if(e.key==='ArrowLeft')show(index-1);if(e.key==='Home')show(0);if(e.key==='End')show(deck.pages.length-1)});setInterval(()=>{const now=performance.now(),delta=(now-last)/1000;last=now;if(!running||document.hidden)return;total+=delta;pageElapsed+=delta;$('time').textContent=String(Math.floor(total/60)).padStart(2,'0')+':'+String(Math.floor(total%60)).padStart(2,'0');if($('auto').checked&&deck.pages[index].seconds>0&&pageElapsed>=deck.pages[index].seconds){if(index<deck.pages.length-1)show(index+1);else{running=false;$('timer').textContent='继续计时'}}},200);show(0);
</script></html>""".replace("DATA</script>", data + "</script>")
    output = Path(output or Path(result["path"]).parent / "rehearsal.html")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html, encoding="utf-8")
    return str(output)
