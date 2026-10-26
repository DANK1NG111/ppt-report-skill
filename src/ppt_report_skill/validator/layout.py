from ppt_report_skill.models import Issue
from ppt_report_skill.layout.capacity import fits
from ppt_report_skill.template.analyzer import walk_shapes
from pptx.enum.shapes import MSO_SHAPE_TYPE


def layout_issues(prs, reference_slides=None):
    issues = []
    w, h = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        # Top-level geometry uses slide coordinates; group children use local coordinates.
        for sh in slide.shapes:
            inherited = False
            if reference_slides and i <= len(reference_slides):
                original = next(
                    (
                        s
                        for s in reference_slides[i - 1].shapes
                        if s.shape_id == sh.shape_id
                    ),
                    None,
                )
                inherited = (
                    original is not None
                    and original.shape_type == sh.shape_type
                    and all(
                        getattr(original, k) == getattr(sh, k)
                        for k in ["left", "top", "width", "height"]
                    )
                )
            if (
                min(sh.left, sh.top) < -9144
                or sh.left + sh.width > w + 9144
                or sh.top + sh.height > h + 9144
            ):
                issues.append(
                    Issue(
                        "TEMPLATE_CROP" if inherited else "OUT_OF_BOUNDS",
                        "warning" if inherited else "error",
                        "模板原有裁切装饰，需结合预览复核。"
                        if inherited
                        else "对象超出页面边界。",
                        i,
                        sh.shape_id,
                    )
                )
            if (
                sh.width <= 0 or sh.height <= 0
            ) and sh.shape_type != MSO_SHAPE_TYPE.LINE:
                issues.append(
                    Issue("INVALID_BOX", "error", "对象尺寸非正数。", i, sh.shape_id)
                )
        for sh in walk_shapes(slide.shapes):
            if sh.has_text_frame and sh.text.strip():
                sizes = [
                    r.font.size.pt
                    for p in sh.text_frame.paragraphs
                    for r in p.runs
                    if r.font.size
                ]
                if sizes and min(sizes) < 12:
                    issues.append(
                        Issue(
                            "SMALL_FONT",
                            "warning",
                            "文字小于 12pt，需检查用途。",
                            i,
                            sh.shape_id,
                        )
                    )
                if sizes and not fits(
                    sh.text, dict(width=sh.width, height=sh.height), max(sizes)
                ):
                    issues.append(
                        Issue(
                            "TEXT_OVERFLOW",
                            "error",
                            "文本估算超出文本框。",
                            i,
                            sh.shape_id,
                        )
                    )
            if sh.has_table:
                if sum(r.height for r in sh.table.rows) > sh.height + 9144:
                    issues.append(
                        Issue(
                            "TABLE_OVERFLOW",
                            "error",
                            "表格行高超出对象。",
                            i,
                            sh.shape_id,
                        )
                    )
    return issues
