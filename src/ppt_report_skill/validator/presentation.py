from dataclasses import asdict
from pathlib import Path
from pptx import Presentation
from ppt_report_skill.models import Issue, save_json
from .structural import structural_issues
from .content import content_issues
from .layout import layout_issues
from .editable import editable_issues


def write_report(issues, output, extra=None):
    errors = sum(i.severity == "error" for i in issues)
    warnings = sum(i.severity == "warning" for i in issues)
    report = {
        "status": "FAIL" if errors else "PASS_WITH_WARNINGS" if warnings else "PASS",
        "errors": errors,
        "warnings": warnings,
        "issues": [asdict(i) for i in issues],
        **(extra or {}),
    }
    if output:
        output = Path(output)
        output.mkdir(parents=True, exist_ok=True)
        save_json(report, output / "qa_report.json")
        lines = [
            "# QA report",
            f"\nStatus: {report['status']}",
            f"\nErrors: {errors}, Warnings: {warnings}",
            "\n检查范围为规则与文件检查，不能替代 PowerPoint 原生打开和逐页视觉复核。",
        ]
        lines += [
            f"\n- {i.severity.upper()} {i.code} / slide {i.slide} / shape {i.shape_id}: {i.message}"
            for i in issues
        ]
        (output / "qa_report.md").write_text("\n".join(lines))
    return report


def validate_presentation(path, output=None, plans=None, request=None, profile=None):
    issues = structural_issues(path)
    if not any(i.code == "INVALID_PPTX" for i in issues):
        try:
            prs = Presentation(path)
        except Exception as exc:
            issues.append(Issue("INVALID_PPTX", "error", str(exc)))
        else:
            references = None
            if profile and plans:
                original = Presentation(profile.template_path)
                layouts = {l.layout_id: l for l in profile.layouts}
                if all(p.layout_id in layouts for p in plans):
                    references = [
                        original.slides[layouts[p.layout_id].source_slide - 1]
                        for p in plans
                    ]
            if plans and len(prs.slides) != len(plans):
                issues.append(Issue("SLIDE_COUNT", "error", "页数与计划不一致。"))
            issues += (
                content_issues(prs, plans, request)
                + layout_issues(prs, references)
                + editable_issues(prs, path, plans)
            )
    return write_report(issues, output)
