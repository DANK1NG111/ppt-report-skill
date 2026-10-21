"""Geometry plus actual rendered-page evidence, with explicit review warnings."""
import hashlib
from pathlib import Path
from statistics import median
import fitz
from PIL import Image,ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_COLOR_TYPE
from ppt_report_skill.models import Issue
from ppt_report_skill.layout.capacity import glyph_units
from .content import compact
from .presentation import write_report


def overlap(a,b):
    area=max(0,min(a.left+a.width,b.left+b.width)-max(a.left,b.left))*max(0,min(a.top+a.height,b.top+b.height)-max(a.top,b.top))
    return area/max(1,min(a.width*a.height,b.width*b.height))


def contrast(a,b):
    def luminance(rgb):
        values=[v/255 for v in rgb]
        linear=[v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4 for v in values]
        return sum(w*v for w,v in zip([.2126,.7152,.0722],linear))
    x,y=sorted([luminance(a),luminance(b)])
    return (y+.05)/(x+.05)


def background_sample(image,shape,width,height):
    # Native filled number markers have their own background inside the frame.
    try:
        if shape.fill.type==1 and shape.fill.fore_color.type==MSO_COLOR_TYPE.RGB:return tuple(shape.fill.fore_color.rgb)
    except (AttributeError,TypeError):pass
    x0=int(shape.left/width*image.width);y0=int(shape.top/height*image.height)
    x1=int((shape.left+shape.width)/width*image.width);y1=int((shape.top+shape.height)/height*image.height)
    points=[(x,y0+1) for x in range(max(0,x0),min(image.width,x1),max(1,(x1-x0)//12))]
    points += [(x0+1,y) for y in range(max(0,y0),min(image.height,y1),max(1,(y1-y0)//12))]
    pixels=[image.getpixel((max(0,min(image.width-1,x)),max(0,min(image.height-1,y))))[:3] for x,y in points]
    return tuple(int(median(p[c] for p in pixels)) for c in range(3)) if pixels else (255,255,255)


def visual_issues(path,rendered,plans=None,profile=None):
    issues=[];prs=Presentation(path);width,height=prs.slide_width,prs.slide_height
    if rendered.get('sha256')!=hashlib.sha256(Path(path).read_bytes()).hexdigest():
        return [Issue('STALE_RENDER','error','渲染回执与当前 PPTX 不一致。')]
    if rendered.get('page_count')!=len(prs.slides) or len(rendered.get('images',[]))!=len(prs.slides):
        return [Issue('RENDER_PAGE_COUNT','error','渲染图片数量与页面数不同。')]
    source=Presentation(profile.template_path) if profile else None
    pdf=fitz.open(rendered['pdf']) if rendered.get('pdf') and Path(rendered['pdf']).is_file() else None
    layouts={l.layout_id:l for l in profile.layouts} if profile else {}
    for i,slide in enumerate(prs.slides,1):
        plan=plans[i-1] if plans and i<=len(plans) else None
        original=None
        if source and plan and plan.layout_id in layouts:original=source.slides[layouts[plan.layout_id].source_slide-1]
        image_path=Path(rendered['images'][i-1])
        if not image_path.is_file():issues.append(Issue('RENDER_MISSING','error','缺少页面渲染图。',i));continue
        with Image.open(image_path) as image:
            image=image.convert('RGB')
            if max(ImageStat.Stat(image).stddev)<1:issues.append(Issue('RENDER_BLANK','error','渲染图近乎纯色。',i))
            text_shapes=[sh for sh in slide.shapes if sh.has_text_frame and sh.text.strip()]
            spans=[span for block in pdf[i-1].get_text('dict')['blocks'] for line in block.get('lines',[]) for span in line['spans']] if pdf and len(pdf)>i-1 else []
            def ink(shape):
                found=[fitz.Rect(s['bbox']) for s in spans if len(compact(s['text']))>=2 and compact(s['text']) in compact(shape.text)]
                if not found:return None
                rect=found[0]
                for other in found[1:]:rect |= other
                return rect
            for index,a in enumerate(text_shapes):
                for b in text_shapes[index+1:]:
                    if overlap(a,b)>.12:
                        ia,ib=ink(a),ink(b)
                        # Empty portions of nominal frames may overlap by design.
                        if ia is None or ib is None or (ia & ib).get_area()>min(ia.get_area(),ib.get_area())*.02:
                            issues.append(Issue('TEXT_OVERLAP','error',f'文字对象与 shape {b.shape_id} 相交。',i,a.shape_id))
                runs=[r for p in a.text_frame.paragraphs for r in p.runs if r.text.strip()]
                rgb=next((tuple(r.font.color.rgb) for r in runs if r.font.color.type==MSO_COLOR_TYPE.RGB),None)
                if rgb and contrast(rgb,background_sample(image,a,width,height))<2.0:
                    issues.append(Issue('LOW_CONTRAST','warning','文字与附近渲染背景对比偏低，需逐页复核。',i,a.shape_id))
            for shape in slide.shapes:
                if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                    iw,ih=shape.image.size
                    effective=iw*(1-shape.crop_left-shape.crop_right)/max(1e-8,ih*(1-shape.crop_top-shape.crop_bottom))
                    ratio=shape.width/max(shape.height,1)
                    if abs(effective/ratio-1)>.035:
                        inherited=original is not None and any(s.shape_id==shape.shape_id and s.shape_type==shape.shape_type and s.left==shape.left and s.top==shape.top and s.width==shape.width and s.height==shape.height for s in original.shapes)
                        issues.append(Issue('TEMPLATE_IMAGE_TRANSFORM' if inherited else 'IMAGE_DISTORTION','warning' if inherited else 'error','图片比例与有效裁剪比例不同。',i,shape.shape_id))
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if any(r.font.size and r.font.size.pt<14 for p in cell.text_frame.paragraphs for r in p.runs):
                                issues.append(Issue('TABLE_READABILITY','warning','表格字号低于 14pt。',i,shape.shape_id))
                if shape.has_chart:
                    if any(len(s.values)>12 for s in shape.chart.series):issues.append(Issue('CHART_READABILITY','warning','单个图表类别超过 12，需核对标签密度。',i,shape.shape_id))
            if plan and plan.slide_type=='content':
                glyphs=sum(glyph_units(s.text) for s in text_shapes)
                if glyphs<35 and plan.visual_type=='text':issues.append(Issue('SPARSE_PAGE','warning','内容较少，需确认是否保留独立页面。',i))
                if glyphs>900:issues.append(Issue('DENSE_PAGE','warning','文字较多，建议拆页或精简。',i))
                title=next((s for s in text_shapes if compact(plan.title)==compact(s.text)),None)
                if title and title.top>height*.25:issues.append(Issue('TITLE_POSITION','warning','正文页标题位置偏低。',i,title.shape_id))
                if original and title and plan.layout_id in layouts and layouts[plan.layout_id].title_box:
                    box=layouts[plan.layout_id].title_box
                    if abs(title.left-box['x'])>91440:issues.append(Issue('TITLE_ALIGNMENT','warning','标题偏离模板原有左边缘。',i,title.shape_id))
            rendered_text=compact(rendered.get('page_texts',['']*len(prs.slides))[i-1])
            if '\ufffd' in rendered_text:issues.append(Issue('MISSING_GLYPH','error','渲染文字含替代字符，可能缺失字体。',i))
            if plan:
                expected=plan.payload.get('text_items',[])
                if plan.visual_type=='table':expected=[str(c) for r in plan.payload['rows'] for c in r]
                if plan.visual_type=='diagram':expected=[n['label'] for n in plan.payload['diagram']['nodes']]
                if plan.slide_type!='contents':expected=[plan.title]+expected
                for value in expected:
                    if compact(value) not in rendered_text:issues.append(Issue('RENDER_TEXT_MISSING','error','渲染未找到计划文字，可能被裁切或缺字：'+value[:45],i))
    if pdf:pdf.close()
    return issues


def validate_visual(path,rendered,output=None,plans=None,profile=None):
    return write_report(visual_issues(path,rendered,plans,profile),output,{'rendered_page_count':rendered.get('page_count'),'manual_review_required':True})
