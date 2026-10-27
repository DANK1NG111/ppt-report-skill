"""User-facing commands with concise errors and retained technical logs."""

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import traceback
from pptx import Presentation
from .models import PresentationRequest, ReportError, SlidePlan, TemplateProfile


def error_details(exc: Exception) -> tuple[str, str]:
    if isinstance(exc, ReportError):
        return exc.code, exc.message
    for kind, code, message in [
        (json.JSONDecodeError, "INVALID_REQUEST", "请求 JSON 格式无效。"),
        (PermissionError, "PERMISSION_DENIED", "文件访问权限不足。"),
        (FileNotFoundError, "INPUT_NOT_FOUND", "输入文件不存在。"),
        (
            (ValueError, TypeError, KeyError),
            "INVALID_INPUT",
            "输入格式或对象数据无效。",
        ),
    ]:
        if isinstance(exc, kind):
            return code, message
    return "INTERNAL_ERROR", "执行失败，请查看技术日志。"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="基于真实模板制作可编辑汇报 PPT")
    parser.add_argument(
        "command",
        choices=[
            "analyze",
            "create",
            "validate",
            "render",
            "regression",
            "finalize",
            "doctor",
            "rehearse",
            "smoke",
        ],
    )
    parser.add_argument(
        "--output", help="输出文件或目录；create 默认 output/final_presentation.pptx"
    )
    parser.add_argument(
        "--template", help="PPTX模板；省略时优先用户模板，否则使用自带简洁模板"
    )
    parser.add_argument("--materials", nargs="*", default=[])
    parser.add_argument("--request", help="PresentationRequest JSON 文件")
    parser.add_argument("--text", help="直接提供的汇报材料")
    parser.add_argument(
        "--input",
        default="output/final_presentation.pptx",
        help="需要校验或渲染的 PPTX",
    )
    parser.add_argument("--build", help="包含逐页方案、请求与模板档案的目录")
    parser.add_argument("--receipt", help="Pipeline 回执 JSON")
    parser.add_argument("--review", help="逐页复核记录 JSON")
    parser.add_argument("--smoke", action="store_true", help="doctor实际生成并渲染自检")
    parser.add_argument("--pdf", action="store_true", help="交付时复制 PDF")
    parser.add_argument(
        "--draft", action="store_true", help="仅生成草稿，跳过完整 QA 流程"
    )
    args = parser.parse_args(argv)
    try:
        return execute(args)
    except Exception as exc:
        code, message = error_details(exc)
        destination = Path(args.output or "output")
        directory = destination.parent if destination.suffix else destination
        try:
            directory.mkdir(parents=True, exist_ok=True)
            with (directory / "error.log").open("a") as log:
                log.write(
                    f"\n{datetime.now(timezone.utc).isoformat()} {code}\n{getattr(exc, 'detail', '')}\n{traceback.format_exc()}"
                )
        except OSError:
            pass
        print(f"{code}: {message}")
        return 2


def execute(args: argparse.Namespace) -> int:
    if args.command == "doctor":
        from .doctor import doctor

        result = doctor(args.output or "output/doctor", smoke=args.smoke)
        print(f"{result['status']}: {result['checks']}")
        return int(result["status"] == "BLOCKED")
    if args.command == "rehearse":
        from .rehearsal import export_rehearsal

        if not args.receipt:
            raise ReportError("INVALID_REQUEST", "rehearse需要--receipt。")
        print(export_rehearsal(args.receipt, args.output))
        return 0
    if args.command in {"create", "analyze"} and not args.template:
        from .starter import default_template

        args.template = str(default_template())
    if args.command == "finalize":
        from .delivery import finalize_delivery

        if not args.receipt or not args.review:
            raise ReportError(
                "INVALID_REQUEST", "finalize 需要 --receipt 和 --review。"
            )
        result = finalize_delivery(args.receipt, args.review, include_pdf=args.pdf)
        print(f"DELIVERED: {result['path']}")
        return 0
    if args.command == "regression":
        from .regression import run_regression

        result = run_regression()
        print(f"{result['status']}: {result['counts']}; {result['elapsed_seconds']}s")
        return int(result["status"] != "PASS")
    if args.command == "render":
        from .renderer.render import render_presentation

        result = render_presentation(args.input, args.output or "output/rendered")
        print(f"PASS: {result['page_count']} rendered; {result['contact_sheet']}")
        return 0
    if args.command == "validate":
        from .validator import validate_presentation

        plans = request = profile = None
        if args.build:
            build = Path(args.build)
            plans = [
                SlidePlan(**x)
                for x in json.loads((build / "slide_plans.json").read_text())
            ]
            if (build / "request.json").exists():
                request = PresentationRequest(
                    **json.loads((build / "request.json").read_text())
                )
            if (build / "template_profile.json").exists():
                profile = TemplateProfile.from_dict(
                    json.loads((build / "template_profile.json").read_text())
                )
        result = validate_presentation(
            args.input, args.output or "output", plans, request, profile
        )
        print(
            f"{result['status']}: errors={result['errors']}, warnings={result['warnings']}"
        )
        return int(bool(result["errors"]))
    if args.command == "create":
        from .draft import create_draft
        from .pipeline import create_presentation

        data = json.loads(Path(args.request).read_text()) if args.request else {}
        output = args.output or "output/final_presentation.pptx"
        if args.draft:
            result = create_draft(
                args.template, args.materials, data, output, args.text
            )
            print(f"DRAFT_UNVALIDATED: {result}")
        else:
            result = create_presentation(
                args.template, args.materials, data, output, args.text
            )
            print(f"{result['status']}: {result['path']}")
        return 0
    if args.command == "analyze":
        from .template.analyzer import analyze_template
        from .renderer.render import render_presentation

        profile_path = Path(
            args.output or "assets/template_profiles/base_template.json"
        )
        if profile_path.suffix != ".json":
            profile_path /= "template_profile.json"
        profile = analyze_template(args.template, profile_path)
        previews = (
            profile_path.parent / "previews" if args.output else Path("assets/previews")
        )
        result = render_presentation(args.template, previews)
        print(f"PASS: {len(profile.layouts)} pages; {result['page_count']} rendered")
        return 0
    if args.command == "smoke":
        path = Path(args.output or "output/smoke.pptx")
        path.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "TEST FIXTURE"
        prs.save(path)
        assert len(Presentation(path).slides) == 1
        print(path)
        return 0
    raise ValueError("未支持的命令")


if __name__ == "__main__":
    raise SystemExit(main())
