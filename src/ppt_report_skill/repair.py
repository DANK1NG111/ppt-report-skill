"""Bounded, evidence-preserving repairs; unresolved issues remain failures."""

from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from ppt_report_skill.layout.capacity import estimated_lines
from ppt_report_skill.models import save_json, Issue
from ppt_report_skill.validator import validate_presentation
from ppt_report_skill.validator.visual import visual_issues
from ppt_report_skill.validator.presentation import write_report
from ppt_report_skill.renderer.render import render_presentation


def repair_once(path, issues, output):
    prs = Presentation(path)
    actions = []
    margin = 182880
    for issue in issues:
        if isinstance(issue, dict):
            issue = Issue(**issue)
        if not issue.slide or not issue.shape_id or issue.slide > len(prs.slides):
            continue
        slide = prs.slides[issue.slide - 1]
        shape = next((s for s in slide.shapes if s.shape_id == issue.shape_id), None)
        if shape is None:
            continue
        before = (shape.left, shape.top, shape.width, shape.height)
        if issue.code == "OUT_OF_BOUNDS":
            if shape.has_table and shape.width > prs.slide_width - 2 * margin:
                ratio = (prs.slide_width - 2 * margin) / shape.width
                for column in shape.table.columns:
                    column.width = max(1, int(column.width * ratio))
                shape.width = sum(c.width for c in shape.table.columns)
            elif shape.width > prs.slide_width - 2 * margin:
                ratio = (prs.slide_width - 2 * margin) / shape.width
                shape.width = int(shape.width * ratio)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    shape.height = int(shape.height * ratio)
            if shape.height <= prs.slide_height - 2 * margin:
                shape.left = max(
                    margin, min(shape.left, prs.slide_width - margin - shape.width)
                )
                shape.top = max(
                    margin, min(shape.top, prs.slide_height - margin - shape.height)
                )
        elif (
            issue.code == "IMAGE_DISTORTION"
            and shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ):
            iw, ih = shape.image.size
            ratio = iw / ih
            target = shape.width / shape.height
            shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
            if ratio > target:
                shape.crop_left = shape.crop_right = (1 - target / ratio) / 2
            else:
                shape.crop_top = shape.crop_bottom = (1 - ratio / target) / 2
            actions.append(
                dict(
                    code=issue.code,
                    slide=issue.slide,
                    shape_id=shape.shape_id,
                    strategy="center_crop_without_stretch",
                )
            )
        elif issue.code == "TEXT_OVERFLOW" and shape.has_text_frame:
            sizes = [
                r.font.size.pt
                for p in shape.text_frame.paragraphs
                for r in p.runs
                if r.font.size
            ]
            size = max(sizes, default=20)
            needed = int(
                estimated_lines(shape.text, max(shape.width - 91440, 1), size)
                * size
                * 1.25
                / 72
                * 914400
                + 91440
            )
            obstacles = [
                s.top
                for s in slide.shapes
                if s != shape
                and s.top > shape.top
                and s.left < shape.left + shape.width
                and s.left + s.width > shape.left
                and (s.has_text_frame and s.text.strip() or s.has_table or s.has_chart)
            ]
            bottom = min(obstacles, default=prs.slide_height - margin) - margin
            if shape.top + needed <= bottom:
                shape.height = needed
        elif issue.code == "TEXT_OVERLAP":
            blockers = [
                s
                for s in slide.shapes
                if s != shape
                and s.has_text_frame
                and s.text.strip()
                and s.left < shape.left + shape.width
                and s.left + s.width > shape.left
                and s.top < shape.top + shape.height
                and s.top + s.height > shape.top
            ]
            if blockers:
                new_top = max(s.top + s.height for s in blockers) + margin
                if new_top + shape.height < prs.slide_height - margin:
                    shape.top = new_top
        after = (shape.left, shape.top, shape.width, shape.height)
        if before != after:
            actions.append(
                dict(
                    code=issue.code,
                    slide=issue.slide,
                    shape_id=shape.shape_id,
                    strategy="geometry_reflow",
                    before=before,
                    after=after,
                )
            )
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return actions


def repair_loop(
    path, output_dir, plans=None, request=None, profile=None, max_iterations=3
):
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    current = Path(path)
    history = []
    for iteration in range(max_iterations + 1):
        structural = validate_presentation(
            current, plans=plans, request=request, profile=profile
        )
        rendered = render_presentation(current, directory / f"render-{iteration}")
        issues = [Issue(**i) for i in structural["issues"]] + visual_issues(
            current, rendered, plans, profile
        )
        report = write_report(issues, directory / f"qa-{iteration}")
        record = {
            "iteration": iteration,
            "path": str(current.resolve()),
            "errors": report["errors"],
            "warnings": report["warnings"],
        }
        history.append(record)
        if not report["errors"] or iteration == max_iterations:
            break
        candidate = directory / f"repaired-{iteration + 1}.pptx"
        actions = repair_once(current, issues, candidate)
        record["actions"] = actions
        if not actions:
            break
        current = candidate
    result = {
        "path": str(current.resolve()),
        "report": report,
        "rendered": rendered,
        "history": history,
        "resolved": not report["errors"],
    }
    save_json(result, directory / "repair_receipt.json")
    return result
