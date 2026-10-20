"""Native text writer retaining template paragraph and run formatting."""
from copy import deepcopy
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import MSO_AUTO_SIZE
from ppt_report_skill.models import ReportError
from ppt_report_skill.layout.capacity import fits


def write_text(shape,text,*,font_size=None,minimum_size=16,allow_small_shrink=True,emphasis=()):
    frame=shape.text_frame
    paragraph=frame.paragraphs[0]
    ppr=deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    runs=[r for p in frame.paragraphs for r in p.runs]
    prototype=next((r for r in runs if r.text.strip()),None)
    rpr=deepcopy(prototype._r.rPr) if prototype is not None and prototype._r.rPr is not None else None
    size=font_size or next((r.font.size.pt for r in runs if r.font.size),20)
    box=dict(width=shape.width,height=shape.height,font_size=size)
    chosen=size
    if not fits(text,box,size):
        lower=max(minimum_size,size*.9)
        if allow_small_shrink and fits(text,box,lower):chosen=lower
        else:raise ReportError('TEXT_OVERFLOW',f'文本超出框容量：{text[:30]}',f'shape={shape.shape_id}, size={size}')
    frame.clear();frame.word_wrap=True;frame.auto_size=MSO_AUTO_SIZE.NONE
    for i,line in enumerate(text.split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        if ppr is not None:p._p.insert(0,deepcopy(ppr))
        p.space_after=Pt(4)
        r=p.add_run()
        if rpr is not None:r._r.insert(0,deepcopy(rpr))
        r.text=line;r.font.size=Pt(chosen)
        if any(word in line for word in emphasis):r.font.bold=True
    return {'shape_id':shape.shape_id,'original_size':size,'font_size':chosen,'shrunk':chosen<size}
