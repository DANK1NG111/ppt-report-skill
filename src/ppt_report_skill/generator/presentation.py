"""Duplicate native source slides while retaining masters, themes and media."""
from copy import deepcopy
from pathlib import Path
import re
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from ppt_report_skill.models import ReportError,save_json
from ppt_report_skill.layout.matcher import match_layout
from ppt_report_skill.template.analyzer import walk_shapes
from .text import write_text
from .images import add_image
from .tables import add_table
from .charts import add_chart
from .shapes import add_diagram

RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def clone_slide(prs,source):
    target=prs.slides.add_slide(source.slide_layout)
    for shape in list(target.shapes):shape._element.getparent().remove(shape._element)
    mapping={}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(('/slideLayout','/notesSlide')):continue
        mapping[rel.rId]=target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,rel.is_external)
    for shape in source.shapes:
        element=deepcopy(shape._element)
        for node in element.iter():
            for key,value in list(node.attrib.items()):
                if key.startswith('{'+RNS+'}') and value in mapping:node.set(key,mapping[value])
        target.shapes._spTree.insert_element_before(element,'p:extLst')
    bg=source._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if bg is not None:
        copied=deepcopy(bg)
        for node in copied.iter():
            for key,value in list(node.attrib.items()):
                if key.startswith('{'+RNS+'}') and value in mapping:node.set(key,mapping[value])
        target._element.cSld.insert(0,copied)
    return target


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def expand_title(shape,text,profile,size):
    # Source templates sometimes use nominal boxes shorter than their font.
    from ppt_report_skill.layout.capacity import estimated_lines
    needed=int(estimated_lines(text,max(shape.width-91440,1),size)*size*1.25/72*914400+91440)
    if needed>shape.height:
        if shape.top+needed>profile.slide_height*.9:raise ReportError('TEXT_OVERFLOW','标题过长，需要较短标题或替换版式。')
        shape.height=needed


def populate(slide,plan,layout,profile,request):
    shapes={s.shape_id:s for s in walk_shapes(slide.shapes)}
    title=shapes.get(layout.title_box['shape_id']) if layout.title_box else None
    bodies=[b for b in layout.body_boxes if b['shape_id'] in shapes]
    receipts=[]
    if title and 'flexible_derived' in layout.secondary_tags:
        for field,attribute in [('x','left'),('y','top'),('width','width'),('height','height')]:setattr(title,attribute,layout.title_box[field])
    if plan.slide_type=='contents':
        slots=sorted(([layout.title_box] if layout.title_box else [])+bodies,key=lambda b:(b['y'],b['x']))
        slots=[b for b in slots if len(b['text'].strip())>1 and not re.search(r'contents|目录',b['text'],re.I)]
        if 'flexible_derived' in layout.secondary_tags:
            slots=[b for b in bodies if b['height']>profile.slide_height*.1]
            if title:write_text(title,'目录',font_size=layout.title_box['font_size'])
        texts=plan.payload.get('text_items',[])
        if len(texts)>len(slots):raise ReportError('LAYOUT_CAPACITY','目录条目超过真实槽位数量。')
        for i,b in enumerate(slots):
            sh=shapes[b['shape_id']];sh.text_frame.clear() if i>=len(texts) else receipts.append(write_text(sh,texts[i],font_size=b['font_size']))
        # Remove unused numbered markers and keep the template's contents artwork.
        numbers=[b for b in bodies if b['text'].strip().isdigit()]
        for i,b in enumerate(sorted(numbers,key=lambda b:b['y'])):
            if i>=len(texts):shapes[b['shape_id']].text_frame.clear()
        return receipts
    if title:
        size=layout.title_box['font_size'];expand_title(title,plan.title,profile,size)
        receipts.append(write_text(title,plan.title,font_size=size))
    elif request.template_mode=='STRICT_TEMPLATE':
        raise ReportError('LAYOUT_NOT_FOUND','选中的模板页没有可编辑标题槽位。')
    if plan.slide_type in {'cover','closing','section'}:
        for b in bodies:shapes[b['shape_id']].text_frame.clear()
        return receipts
    items=plan.payload.get('text_items',[])
    # Existing long body areas are preferred to small keyword labels.
    areas=sorted([b for b in bodies if b['height']>profile.slide_height*.10 and b['width']>profile.slide_width*.15],key=lambda b:(b['y'],b['x']))
    if not areas:areas=sorted(bodies,key=lambda b:b['width']*b['height'],reverse=True)
    for b in bodies:shapes[b['shape_id']].text_frame.clear()
    for b in layout.image_boxes:
        if b['shape_id'] in shapes:remove_shape(shapes[b['shape_id']])
    for b in layout.chart_boxes+layout.table_boxes:
        if b['shape_id'] in shapes:remove_shape(shapes[b['shape_id']])
    if plan.visual_type=='text':
        if len(items)>len(areas):raise ReportError('LAYOUT_CAPACITY',f'页面有 {len(items)} 条内容，模板仅有 {len(areas)} 个正文区域。')
        for text,b in zip(items,areas):
            sh=shapes[b['shape_id']]
            receipts.append(write_text(sh,text,font_size=max(16,b['font_size'])))
        # Replace preceding native subheadings with a neutral, source-derived label.
        for text,b in zip(items,areas):
            labels=[x for x in bodies if x not in areas and x['y']<=b['y'] and b['y']-x['y']<914400*.8]
            if labels:
                label=max(labels,key=lambda x:x['y'])
                candidate=text.split('：')[0] if not text.startswith('[待补充') and '：' in text and len(text.split('：')[0])<=18 else ''
                if candidate:receipts.append(write_text(shapes[label['shape_id']],candidate,font_size=label['font_size']))
    else:
        if not areas:raise ReportError('LAYOUT_NOT_FOUND','模板没有可放置内容的区域。')
        # Use the union of existing body areas; fixed objects remain untouched.
        x=min(b['x'] for b in areas);y=min(b['y'] for b in areas)
        box={'x':x,'y':y,'width':max(b['x']+b['width'] for b in areas)-x,'height':max(b['y']+b['height'] for b in areas)-y}
        # A chart/table owns the body region, so unused native text objects are removed.
        for b in bodies:
            if b['shape_id'] in shapes:remove_shape(shapes[b['shape_id']])
        accent=next((c for c in profile.colors if c.upper() not in {'FFFFFF','000000'}),'A6383B')
        font=next((f for f in profile.fonts if '微' in f),profile.fonts[0] if profile.fonts else 'Arial')
        if plan.visual_type=='table':add_table(slide,plan.payload['rows'],box,accent=accent,font_name=font)
        elif plan.visual_type=='chart':add_chart(slide,plan.payload['chart'],box,accent=accent,font_name=font)
        elif plan.visual_type=='diagram':add_diagram(slide,plan.payload['diagram'],box,accent=accent,font_name=font)
        elif plan.visual_type=='image':add_image(slide,plan.payload['image'],layout.image_boxes[0] if layout.image_boxes else box,plan.payload.get('fit','contain'))
        else:raise ReportError('UNSUPPORTED_VISUAL',f'不支持的 visual_type：{plan.visual_type}')
    if request.speaker_notes or plan.source_refs:
        lines=[f'讲述时间建议：{plan.estimated_speaking_time:.1f} 秒',plan.core_message,'来源：']
        lines += [f"{r['source']} / {r['location']} / {r['fact_id']}\n{r.get('original_text','')}" for r in plan.source_refs]
        slide.notes_slide.notes_text_frame.text='\n'.join(lines)
    return receipts


def generate_presentation(template,profile,plans,request,output,build_dir=None):
    prs=Presentation(template);original=list(prs.slides);receipts=[]
    for plan in plans:
        layout=next((l for l in profile.layouts if l.layout_id==plan.layout_id),None) or match_layout(plan,profile,request.template_mode)
        slide=clone_slide(prs,original[layout.source_slide-1])
        receipts.append({'slide':plan.slide_number,'layout_id':layout.layout_id,'source_slide':layout.source_slide,'text':populate(slide,plan,layout,profile,request)})
    for slide_id in list(prs.slides._sldIdLst)[:len(original)]:
        prs.part.drop_rel(slide_id.rId);prs.slides._sldIdLst.remove(slide_id)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);prs.save(output)
    if len(Presentation(output).slides)!=len(plans):raise ReportError('GENERATION_FAILED','生成后页数与计划不一致。')
    if build_dir:
        save_json(receipts,Path(build_dir)/'generation_receipt.json')
        save_json([__import__('dataclasses').asdict(p) for p in plans],Path(build_dir)/'slide_plans.json')
    return output
