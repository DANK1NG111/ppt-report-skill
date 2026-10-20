from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE,MSO_CONNECTOR_TYPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.util import Pt
from pptx.dml.color import RGBColor
from ppt_report_skill.models import ReportError
from .text import write_text


def add_diagram(slide,data,box,accent='A6383B',font_name=None):
    kind=data.get('kind','process');nodes=data.get('nodes',[])
    if kind not in {'process','timeline','stages','organization','architecture'} or not nodes:raise ReportError('DIAGRAM_INVALID','流程图类型或节点无效。')
    ids=[n['id'] for n in nodes]
    if len(set(ids))!=len(ids):raise ReportError('DIAGRAM_INVALID','节点 ID 重复。')
    edges=data.get('edges',[[ids[i],ids[i+1]] for i in range(len(ids)-1)])
    if any(a not in ids or b not in ids for a,b in edges):raise ReportError('DIAGRAM_INVALID','连接引用不存在的节点。')
    levels={n:0 for n in ids}
    if kind in {'organization','architecture'}:
        for iteration in range(len(ids)):
            changed=False
            for a,b in edges:
                if levels[b]<=levels[a]:levels[b]=levels[a]+1;changed=True
            if not changed:break
        else:raise ReportError('DIAGRAM_INVALID','层次图存在循环。')
    else:levels={n:0 for n in ids}
    groups={level:[n for n in nodes if levels[n['id']]==level] for level in sorted(set(levels.values()))}
    shapes={};max_level=max(groups)
    for level,items in groups.items():
        cell_w=box['width']/len(items);cell_h=box['height']/(max_level+1)
        for j,node in enumerate(items):
            w=int(cell_w*.77);h=int(min(cell_h*.5,914400*1.1))
            x=int(box['x']+j*cell_w+(cell_w-w)/2);y=int(box['y']+level*cell_h+(cell_h-h)/2)
            sh=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,x,y,w,h)
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(accent)
            sh.line.color.rgb=RGBColor.from_string(accent)
            sh.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
            write_text(sh,node['label'],font_size=18)
            for p in sh.text_frame.paragraphs:
                p.alignment=PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.color.rgb=RGBColor(255,255,255)
                    if font_name:r.font.name=font_name
            shapes[node['id']]=sh
    connectors=[]
    for a,b in edges:
        s,t=shapes[a],shapes[b]
        connection=slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT,0,0,1,1)
        connection.begin_connect(s,3 if levels[a]==levels[b] else 2)
        connection.end_connect(t,1 if levels[a]==levels[b] else 0)
        connection.line.color.rgb=RGBColor.from_string(accent);connectors.append(connection)
    return list(shapes.values())+connectors
