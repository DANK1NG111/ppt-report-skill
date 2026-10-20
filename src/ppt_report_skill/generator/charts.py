import math
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.util import Pt
from pptx.dml.color import RGBColor
from ppt_report_skill.models import ReportError

KINDS={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE,'pie':XL_CHART_TYPE.PIE}


def add_chart(slide,data,box,accent='A6383B',font_name=None):
    kind=data.get('kind','column');categories=data.get('categories',[]);series=data.get('series',[])
    if kind not in KINDS or not categories or not series:raise ReportError('CHART_INVALID','图表类型或数据无效。')
    if not data.get('unit') or not data.get('time'):raise ReportError('CHART_INVALID','图表须明确单位与时间范围。')
    if kind=='pie' and (len(categories)>6 or len(series)!=1):raise ReportError('CHART_INVALID','饼图仅支持最多六类、单一系列。')
    names=[s.get('name') for s in series]
    if any(not n for n in names) or len(set(names))!=len(names):raise ReportError('CHART_INVALID','图表系列须有唯一名称。')
    chart_data=CategoryChartData();chart_data.categories=[str(c) for c in categories]
    for item in series:
        values=item['values']
        if len(values)!=len(categories) or any(not isinstance(v,(int,float)) or not math.isfinite(v) for v in values):raise ReportError('CHART_INVALID','系列长度与类别不一致或数值无效。')
        if kind=='pie' and (any(v<0 for v in values) or not sum(values)):raise ReportError('CHART_INVALID','饼图数据须非负且合计大于零。')
        chart_data.add_series(item['name'],values)
    shape=slide.shapes.add_chart(KINDS[kind],box['x'],box['y'],box['width'],box['height'],chart_data)
    chart=shape.chart;chart.has_title=True
    chart.chart_title.text_frame.text=f"{data.get('title','数据')}（{data['unit']}，{data['time']}）"
    chart.font.size=Pt(16)
    if font_name:chart.font.name=font_name
    for p in chart.chart_title.text_frame.paragraphs:
        for r in p.runs:r.font.size=Pt(18)
    chart.has_legend=kind=='pie' or len(series)>1
    if chart.has_legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
    if kind!='pie':
        chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string(accent)
    return shape
