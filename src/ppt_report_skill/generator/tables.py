from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from ppt_report_skill.layout.capacity import glyph_units, estimated_lines
from ppt_report_skill.models import ReportError


def table_pages(rows, max_rows=8, max_cols=5):
    if not rows or not rows[0]:
        raise ReportError("TABLE_EMPTY", "表格为空。")
    count = len(rows[0])
    if any(len(r) != count for r in rows):
        raise ReportError("TABLE_INVALID", "表格每行列数必须一致。")
    panels = []
    indices = (
        [list(range(count))]
        if count <= max_cols
        else [
            [0] + list(range(start, min(start + max_cols - 1, count)))
            for start in range(1, count, max_cols - 1)
        ]
    )
    for columns in indices:
        for start in range(1, max(len(rows), 2), max_rows - 1):
            panels.append(
                [[str(rows[0][i]) for i in columns]]
                + [
                    [str(row[i]) for i in columns]
                    for row in rows[start : start + max_rows - 1]
                ]
            )
    return panels


def add_table(
    slide, rows, box, font_size=18, accent="A6383B", font_name=None, highlight=()
):
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ReportError("TABLE_INVALID", "表格数据不完整。")
    values = [[str(x) for x in row] for row in rows]
    weights = [
        max(4, min(30, max(glyph_units(row[i]) for row in values)))
        for i in range(len(values[0]))
    ]
    widths = [int(box["width"] * w / sum(weights)) for w in weights]
    heights = [
        int(
            max(
                estimated_lines(text, max(width - 91440, 1), font_size)
                for text, width in zip(row, widths)
            )
            * font_size
            * 1.35
            / 72
            * 914400
            + 91440
        )
        for row in values
    ]
    if sum(heights) > box["height"]:
        raise ReportError("TABLE_OVERFLOW", "表格内容过高，请拆表或拆页。")
    shape = slide.shapes.add_table(
        len(values), len(values[0]), box["x"], box["y"], box["width"], sum(heights)
    )
    table = shape.table
    for col, width in zip(table.columns, widths):
        col.width = width
    for i, (row, height) in enumerate(zip(table.rows, heights)):
        row.height = height
        for j, cell in enumerate(row.cells):
            cell.text = values[i][j]
            cell.margin_left = cell.margin_right = 45720
            for p in cell.text_frame.paragraphs:
                p.alignment = (
                    PP_ALIGN.RIGHT
                    if i and values[i][j].replace(".", "").replace("%", "").isdigit()
                    else PP_ALIGN.LEFT
                )
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    if font_name:
                        run.font.name = font_name
                    run.font.bold = i == 0 or (i, j) in highlight
                    run.font.color.rgb = RGBColor.from_string(
                        "FFFFFF" if i == 0 else "333333"
                    )
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                accent if i == 0 else "FFFFFF"
            )
    return shape
