"""Create our own neutral starter; no third-party template assets are bundled."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def create_starter(path):
    path = Path(path)
    if path.exists():
        return path
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    def text(slide, value, x, y, w, h, size, color="17365D"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text = value
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(size)
                r.font.color.rgb = RGBColor.from_string(color)

    for kind in ["cover", "contents", "section", "content", "closing"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F8FC")
        if kind in {"cover", "section", "closing"}:
            text(
                slide,
                {"cover": "汇报主题", "section": "章节标题", "closing": "感谢聆听"}[
                    kind
                ],
                0.8,
                2.6,
                11.7,
                1.5,
                44,
            )
        elif kind == "contents":
            text(slide, "目录", 0.8, 0.4, 11.7, 0.8, 32)
            for i in range(4):
                text(slide, f"目录条目{i + 1}", 1.05, 1.75 + i * 1.05, 10.8, 0.65, 24)
        else:
            text(slide, "正文标题", 0.8, 0.4, 11.7, 0.8, 32)
            for i in range(3):
                text(
                    slide,
                    "正文示例，替换为材料中的事实与证据。",
                    0.9,
                    1.8 + i * 1.55,
                    11.5,
                    1.1,
                    22,
                    "202020",
                )
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def default_template(root=None):
    root = Path(root) if root else Path(__file__).resolve().parents[2]
    original = root / "assets/templates/base_template.pptx"
    return (
        original
        if original.exists()
        else create_starter(root / "assets/templates/starter_template.pptx")
    )
