from collections import Counter
import hashlib
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from ppt_report_skill.models import LayoutProfile, TemplateProfile, ReportError, save_json
from .classifier import classify

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def box(shape):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs] if shape.has_text_frame else []
    sizes = [r.font.size.pt for r in runs if r.font.size]
    return dict(shape_id=shape.shape_id, name=shape.name, x=shape.left, y=shape.top, width=shape.width, height=shape.height,
                text=shape.text if shape.has_text_frame else '', font_size=max(sizes, default=20),
                fonts=sorted({r.font.name for r in runs if r.font.name}), kind=str(shape.shape_type))


def analyze_template(path, output=None):
    path = Path(path).resolve()
    if not path.is_file():
        raise ReportError("TEMPLATE_NOT_FOUND", f"模板不存在：{path}")
    try:
        prs = Presentation(path)
    except Exception as exc:
        raise ReportError("INVALID_PPTX", "模板无法读取。", str(exc)) from exc
    width, height = prs.slide_width, prs.slide_height
    fonts, colors, themes, theme_fonts = Counter(), Counter(), {}, set()
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith('.xml') and (name.startswith('ppt/slides/') or name.startswith('ppt/slideMasters/') or name.startswith('ppt/theme/')):
                root = etree.fromstring(archive.read(name))
                colors.update(root.xpath('//a:srgbClr/@val', namespaces=NS))
                fonts.update(root.xpath('//a:latin/@typeface | //a:ea/@typeface', namespaces=NS))
                if name.startswith('ppt/theme/'):
                    theme_fonts.update(root.xpath('//a:fontScheme//a:latin/@typeface | //a:fontScheme//a:ea/@typeface | //a:fontScheme//a:font/@typeface', namespaces=NS))
                    for node in root.xpath('//a:clrScheme/*', namespaces=NS):
                        themes[f'{name}:{etree.QName(node).localname}'] = node[0].get('val') or node[0].get('lastClr')
    image_counts = Counter()
    for slide in prs.slides:
        for shape in walk_shapes(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_counts[shape.image.sha1] += 1
    layouts, examples, fixed, logos, headers, footers, placeholders, backgrounds = [], [], [], [], [], [], [], []
    editable = Counter()
    for index, slide in enumerate(prs.slides,1):
        shapes = list(walk_shapes(slide.shapes))
        texts = [box(s) for s in shapes if s.has_text_frame and s.text.strip()]
        images = [box(s) for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        tables = [box(s) for s in shapes if s.has_table]
        charts = [box(s) for s in shapes if s.has_chart]
        kind, tags = classify('\n'.join(b['text'] for b in texts), texts, images, width,height,index,len(prs.slides),len(tables),len(charts))
        title = next((box(s) for s in shapes if s.is_placeholder and s.placeholder_format.type in (1,3) and s.has_text_frame and s.text.strip()),None)
        if title is None and texts:
            title = max(texts, key=lambda b:b['font_size']*min(b['width']/width*3,1))
        slide_fixed, content_images = [], []
        for shape in shapes:
            b = box(shape)
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            # Border images and large backgrounds are conservatively protected.
            border = shape.top+shape.height < height*.15 or shape.top > height*.91
            large = shape.width*shape.height > width*height*.65
            is_logo = is_picture and border and shape.width < width*.5
            repeated = is_picture and image_counts[shape.image.sha1] >= 3 and border
            is_fixed_text = shape.has_text_frame and bool(re.search(r'出品|版权所有|copyright',shape.text,re.I))
            if is_logo or repeated or (is_picture and large) or is_fixed_text:
                b['reason'] = 'logo_candidate' if is_logo else 'background' if large else 'repeated_border_or_attribution'
                b['slide'] = index
                slide_fixed.append(b);fixed.append(b)
                if is_logo: logos.append(b)
                if border and shape.top < height*.15: headers.append(b)
                if shape.top > height*.85 or is_fixed_text: footers.append(b)
            elif is_picture:
                content_images.append(b)
            if shape.is_placeholder:
                placeholders.append({**b,'slide':index,'placeholder_type':str(shape.placeholder_format.type)})
            editable['text' if shape.has_text_frame else 'table' if shape.has_table else 'chart' if shape.has_chart else 'picture' if is_picture else 'shape'] += 1
        protected_ids={b['shape_id'] for b in slide_fixed}
        bodies=[b for b in texts if b != title and b['shape_id'] not in protected_ids]
        # Capacity uses real text geometry with a conservative Chinese glyph estimate.
        capacities=[int(max(0,b['width']/914400/(max(b['font_size'],16)/72)*.85)*max(1,b['height']/914400/(max(b['font_size'],16)/72*1.35))) for b in bodies]
        recommended=sum(capacities)
        lp=LayoutProfile(f'slide-{index}',slide.slide_layout.name,kind,index,title_box=title,body_boxes=bodies,image_boxes=content_images,chart_boxes=charts,table_boxes=tables,fixed_shapes=slide_fixed,recommended_content_size=recommended,max_content_size=int(recommended*1.2),visual_density=min(1,sum(b['width']*b['height'] for b in texts)/(width*height)),secondary_tags=tags,content_shapes=bodies+content_images+charts+tables)
        layouts.append(lp)
        examples.append({'slide_number':index,'primary_type':kind,'secondary_tags':tags,'texts':[b['text'] for b in texts],'shape_count':len(shapes)})
        backgrounds.append({'slide':index,'explicit':slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld/{http://schemas.openxmlformats.org/presentationml/2006/main}bg') is not None,'inherits_from_layout':True})
    profile=TemplateProfile(str(path),hashlib.sha256(path.read_bytes()).hexdigest(),width,height,width/height,
        [{'name':m.name,'layouts':[l.name for l in m.slide_layouts], 'shape_count':len(m.shapes)} for m in prs.slide_masters],layouts,
        [f for f,_ in fonts.most_common() if f and not f.startswith('+')],sorted(theme_fonts),[c for c,_ in colors.most_common()],themes,backgrounds,fixed,logos,headers,footers,
        [p for p in placeholders if 'SLIDE_NUMBER' in p['placeholder_type']],placeholders,examples,
        {l.layout_id:{'recommended':l.recommended_content_size,'maximum':l.max_content_size} for l in layouts},dict(editable),
        {'微软雅黑':['Microsoft YaHei','Noto Sans CJK SC','PingFang SC'],'思源宋体':['Source Han Serif SC','Songti SC'],'default':['Arial','Noto Sans CJK SC','PingFang SC']})
    if output:
        output=Path(output);save_json(profile,output)
        lines=['# 模板档案',f'\n比例：{profile.aspect_ratio:.4f}，{len(prs.slides)} 页，{len(profile.masters)} 个母版。',f'\n主要字体：{", ".join(profile.fonts[:8])}',f'\n常用显式色：{", ".join(profile.colors[:10])}', '\n样式来自真实模板主题、母版及各页对象。固定元素为几何和重复性启发式候选，使用前需查看预览。', '\n## 页面用途和容量','\n|来源页|类型|推荐字符量|固定对象数|','|---|---|---|---|']
        lines += [f'|{l.source_slide}|{l.layout_type}|{l.recommended_content_size}|{len(l.fixed_shapes)}|' for l in layouts]
        lines += ['\n## 使用约束','保留母版、背景、Logo、边缘图片和署名。unknown 说明页不得用于生成。正文超载优先拆页，不以极小字号兜底。字体缺失需记录替代。容量是估算值，必须结合渲染检查。','模板中的版权说明属于原作者条款，生成内容不应把学校介绍等示例材料当作用户事实。']
        output.with_suffix('.md').write_text('\n'.join(lines),encoding='utf-8')
    return profile
