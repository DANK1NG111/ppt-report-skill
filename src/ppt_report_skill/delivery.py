"""Hash-bound delivery gate after an agent or person reviews actual rendered pages."""

import hashlib
import json
from pathlib import Path
import shutil
from pptx import Presentation
from .models import ReportError, save_json
from .validator.content import slide_text


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def warning_key(issue):
    return f"{issue['code']}:{issue.get('slide')}:{issue.get('shape_id')}"


def finalize_delivery(receipt, review, *, include_pdf=False):
    """Do not create a review automatically: its observations must come from inspection."""
    result = (
        json.loads(Path(receipt).read_text())
        if isinstance(receipt, (str, Path))
        else receipt
    )
    checked = (
        json.loads(Path(review).read_text())
        if isinstance(review, (str, Path))
        else review
    )
    path = Path(result["path"])
    rendered = result["rendered"]

    def reject(message):
        raise ReportError("DELIVERY_NOT_READY", message)

    if result["qa"]["errors"]:
        reject("仍有 QA Error，不能交付。")
    if digest(path) != result["sha256"] or checked.get("sha256") != result["sha256"]:
        reject("复核记录与当前 PPTX 不匹配。")
    if rendered["sha256"] != result["sha256"]:
        reject("渲染与当前 PPTX 不匹配。")
    artifacts = rendered.get("artifact_hashes", {})
    expected = [rendered["pdf"], rendered["contact_sheet"], *rendered["images"]]
    if not artifacts or any(
        not Path(p).is_file() or artifacts.get(p) != digest(p) for p in expected
    ):
        reject("渲染产物缺失或已改变，请重新渲染和复核。")
    if (
        len(Presentation(path).slides) != result["slide_count"]
        or rendered["page_count"] != result["slide_count"]
    ):
        reject("页数不一致。")
    pages = checked.get("pages", [])
    if sorted(p.get("slide") for p in pages) != list(
        range(1, result["slide_count"] + 1)
    ) or any(p.get("status") != "PASS" or not p.get("observation") for p in pages):
        reject("必须逐页记录视觉复核结果。")
    if not checked.get("facts_checked") or not checked.get("section_order_checked"):
        reject("必须核对来源事实和章节顺序。")
    native = checked.get("native_open", {})
    if (
        native.get("application") != "Microsoft PowerPoint"
        or native.get("repair_prompt") is not False
        or not native.get("observation")
    ):
        reject("缺少 PowerPoint 无修复提示的原生打开证据。")
    dispositions = checked.get("warning_dispositions", {})
    for issue in result["qa"]["issues"]:
        if issue["severity"] == "warning" and not dispositions.get(warning_key(issue)):
            reject("尚有 Warning 未记录处置。")
    placeholders = [
        i
        for i, s in enumerate(Presentation(path).slides, 1)
        if "[待补充" in slide_text(s)
    ]
    output = path.parent
    pdf = None
    if include_pdf:
        pdf = output / "final_presentation.pdf"
        shutil.copy2(rendered["pdf"], pdf)
    save_json(checked, output / "visual_review.json")
    warnings = [
        f"- {warning_key(i)}：{dispositions[warning_key(i)]}"
        for i in result["qa"]["issues"]
        if i["severity"] == "warning"
    ]
    lines = [
        "# 生成报告",
        f"\n状态：DELIVERED\n\n文件：{path.name}\n\n页数：{result['slide_count']}\n\n使用模板：{result['template']}\n\n汇报类型：{result['presentation_type']}",
        "\n材料来源：",
        *["- " + s for s in result["sources"]],
        f"\n待补充内容：{'第 ' + ', '.join(map(str, placeholders)) + ' 页' if placeholders else '无'}",
        "\n字体兼容：模板字体声明保留；本机缺少 "
        + ("、".join(result.get("fonts", {}).get("missing_families", [])) or "无")
        + "。",
        *rendered.get("font_substitutions", []),
        "\n兼容性：PowerPoint 原生打开无修复提示。PDF/PNG 使用 LibreOffice 渲染，字体和换行可能与另一台电脑略有不同。",
        f"\nQA：Error 0；Warning {result['qa']['warnings']}，已逐项复核；全部页面已渲染及查看。",
        *warnings,
        "\n推断与默认值：",
        *["- " + s for s in result.get("assumptions", [])],
        f"\n可追溯记录：{result['build']}\n\nPPTX SHA256：{result['sha256']}",
    ]
    (output / "generation_report.md").write_text("\n".join(lines))
    delivered = {
        **result,
        "status": "DELIVERED",
        "pdf": str(pdf) if pdf else None,
        "generation_report": str(output / "generation_report.md"),
        "visual_review": str(output / "visual_review.json"),
    }
    save_json(delivered, output / "delivery_receipt.json")
    return delivered
