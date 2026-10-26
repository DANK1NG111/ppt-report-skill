import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.generator.images import add_image


@pytest.mark.parametrize(
    "size", [(800, 400), (400, 800), (500, 500), (5000, 3000), (20, 15)]
)
@pytest.mark.parametrize("fit", ["contain", "center_crop"])
def test_image_aspect(size, fit, tmp_path):
    f = tmp_path / "image.png"
    Image.new("RGB", size, "#a63d40").save(f)
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    box = dict(x=Inches(1), y=Inches(1), width=Inches(4), height=Inches(3))
    pic = add_image(s, f, box, fit)
    effective = (pic.image.size[0] * (1 - pic.crop_left - pic.crop_right)) / (
        pic.image.size[1] * (1 - pic.crop_top - pic.crop_bottom)
    )
    assert abs(effective - pic.width / pic.height) < 0.001
    assert box["x"] <= pic.left and pic.left + pic.width <= box["x"] + box["width"] + 1
