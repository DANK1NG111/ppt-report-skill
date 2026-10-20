import pytest
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.generator.tables import add_table,table_pages


@pytest.mark.parametrize('rows,cols,long',[(3,3,False),(5,5,False),(10,5,False),(5,3,True),(5,5,False)])
def test_native_tables(rows,cols,long,tmp_path):
    data=[[f'列{j}' for j in range(cols)]]+[[('长文本测试需要换行' if long else str(i*cols+j)) for j in range(cols)] for i in range(rows-1)]
    p=Presentation();pages=table_pages(data,max_rows=6)
    for panel in pages:
        s=p.slides.add_slide(p.slide_layouts[6]);add_table(s,panel,dict(x=Inches(.5),y=Inches(.5),width=Inches(9),height=Inches(6)),font_size=18)
    path=tmp_path/'tables.pptx';p.save(path);read=Presentation(path)
    assert all(s.shapes[0].has_table for s in read.slides)
    actual=[r.cells[0].text for s in read.slides for r in list(s.shapes[0].table.rows)[1:]]
    assert actual==[r[0] for r in data[1:]]
