import pytest
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.generator.shapes import add_diagram


@pytest.mark.parametrize('kind',['process','timeline','stages','organization','architecture'])
def test_editable_diagrams(kind,tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);data={'kind':kind,'nodes':[{'id':'a','label':'收集材料'},{'id':'b','label':'核对事实'},{'id':'c','label':'交付成果'}]}
    shapes=add_diagram(s,data,dict(x=Inches(.5),y=Inches(.5),width=Inches(9),height=Inches(5)))
    assert len(shapes)==5
    path=tmp_path/'diagram.pptx';p.save(path);s=Presentation(path).slides[0]
    assert sum(sh._element.tag.endswith('cxnSp') for sh in s.shapes)==2
    assert {sh.text for sh in s.shapes if sh.has_text_frame}=={'收集材料','核对事实','交付成果'}
