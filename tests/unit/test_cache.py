from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
import ppt_report_skill.template.profile as module


def test_profile_cache_hit_change_and_corruption(tmp_path,monkeypatch):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[0]);s.shapes.title.text='TEST FIXTURE'
    source=tmp_path/'source.pptx';p.save(source);original=module.analyze_template;calls=[]
    def spy(*args,**kwargs):calls.append(1);return original(*args,**kwargs)
    monkeypatch.setattr(module,'analyze_template',spy)
    one,a=module.cached_profile(source,tmp_path/'cache');two,b=module.cached_profile(source,tmp_path/'cache')
    assert a['hit'] is False and b['hit'] is True and len(calls)==1
    assert one.sha256==two.sha256
    with ZipFile(source,'a') as archive:archive.comment=b'Changed test fixture bytes'
    three,c=module.cached_profile(source,tmp_path/'cache')
    assert c['hit'] is False and three.sha256!=one.sha256 and len(calls)==2
    Path(c['cache_path']).write_text('{bad json')
    module.cached_profile(source,tmp_path/'cache');assert len(calls)==3
