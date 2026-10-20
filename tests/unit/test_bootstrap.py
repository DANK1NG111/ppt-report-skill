from ppt_report_skill.cli import main
from pptx import Presentation


def test_pptx_roundtrip(tmp_path):
    path = tmp_path / "smoke.pptx"
    assert main(["smoke", "--output", str(path)]) == 0
    assert Presentation(path).slides[0].shapes.title.text == "TEST FIXTURE"
