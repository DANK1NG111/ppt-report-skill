from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches,Pt
from ppt_report_skill.validator import validate_presentation
from ppt_report_skill.validator.structural import structural_issues


def test_bad_geometry_and_template_residue(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Inches(-1),0,Inches(2),Inches(.2));sh.text='单击此处添加标题'+('过量内容'*100)
    sh.text_frame.paragraphs[0].runs[0].font.size=Pt(24)
    path=tmp_path/'bad.pptx';p.save(path);r=validate_presentation(path,tmp_path/'qa')
    assert {'OUT_OF_BOUNDS','TEXT_OVERFLOW','TEMPLATE_SAMPLE'}<={i['code'] for i in r['issues']}
    assert r['errors']>=3 and (tmp_path/'qa/qa_report.md').exists()


def test_broken_relationship(tmp_path):
    p=tmp_path/'bad.pptx'
    with ZipFile(p,'w') as z:
        z.writestr('[Content_Types].xml','<Types/>');z.writestr('ppt/presentation.xml','<presentation/>')
        z.writestr('_rels/.rels','<Relationships><Relationship Id="rId1" Target="missing.xml"/></Relationships>')
    assert 'BROKEN_RELATIONSHIP' in {i.code for i in structural_issues(p)}


def test_valid_native_file(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2));sh.text='TEST FIXTURE 原生标题'
    out=tmp_path/'good.pptx';p.save(out);assert validate_presentation(out)['errors']==0


def test_inherited_crop_is_reviewed_not_silently_ignored(tmp_path):
    from ppt_report_skill.models import SlidePlan
    from ppt_report_skill.template.analyzer import analyze_template
    from ppt_report_skill.generator.presentation import generate_presentation
    from ppt_report_skill.models import PresentationRequest
    p=analyze_template('assets/templates/base_template.pptx')
    plans=[SlidePlan(1,'封面','cover','TEST FIXTURE','TEST FIXTURE',[],[],layout_requirement='cover')]
    out=generate_presentation(p.template_path,p,plans,PresentationRequest(),tmp_path/'crop.pptx')
    r=validate_presentation(out,plans=plans,profile=p)
    assert r['errors']==0 and any(i['code']=='TEMPLATE_CROP' for i in r['issues'])
