import pytest
from pptx import Presentation
from pptx.util import Inches,Pt
from ppt_report_skill.generator.text import write_text
from ppt_report_skill.models import ReportError


@pytest.mark.parametrize('text',['中文正文保持可编辑。','English report with evidence.','项目 Project 完成率 87.6%','2026-10-01 / 12.50万元','第一点\n第二点\n第三点','这是一个长度适中的中文信息型标题'])
def test_text_native(text,tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(3));r=sh.text_frame.paragraphs[0].add_run();r.text='prototype';r.font.name='Arial';r.font.size=Pt(24)
    receipt=write_text(sh,text)
    out=tmp_path/'text.pptx';p.save(out)
    assert Presentation(out).slides[0].shapes[0].text==text
    assert receipt['font_size']>=16


def test_overflow_rejected_instead_of_tiny_font():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(0,0,Inches(1),Inches(.3))
    with pytest.raises(ReportError,match='TEXT_OVERFLOW'):write_text(sh,'超长内容'*200,font_size=24)
