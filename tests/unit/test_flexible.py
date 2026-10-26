from pptx import Presentation
from pptx.util import Inches, Pt
from ppt_report_skill.template.analyzer import analyze_template
from ppt_report_skill.layout.flexible import extend_layouts
from ppt_report_skill.models import SlidePlan, PresentationRequest
from ppt_report_skill.generator.presentation import generate_presentation


def test_flexible_derives_missing_layout_without_template_specific_logic(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    for text, y, h in [("TEST FIXTURE 标题", 0.3, 1), ("原生正文区域内容" * 5, 2, 3)]:
        sh = s.shapes.add_textbox(Inches(1), Inches(y), Inches(8), Inches(h))
        sh.text = text
        sh.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    path = tmp_path / "template.pptx"
    p.save(path)
    profile = analyze_template(path)
    assert "cover" not in {l.layout_type for l in profile.layouts}
    assert len(extend_layouts(profile, "STRICT_TEMPLATE").layouts) == 1
    extend_layouts(profile, "FLEXIBLE_TEMPLATE")
    plan = SlidePlan(
        1,
        "封面",
        "cover",
        "TEST FIXTURE",
        "TEST FIXTURE",
        [],
        [],
        layout_requirement="cover",
    )
    out = generate_presentation(
        path,
        profile,
        [plan],
        PresentationRequest(template_mode="FLEXIBLE_TEMPLATE"),
        tmp_path / "out.pptx",
    )
    assert Presentation(out).slides[0].shapes[0].top > Inches(1)
