from zipfile import ZipFile
import pytest
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.generator.charts import add_chart


@pytest.mark.parametrize('kind',['column','bar','line','pie'])
def test_native_chart_workbook(kind,tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);data=dict(kind=kind,title='TEST FIXTURE',unit='个',time='2026年第三季度',categories=['甲','乙','丙'],series=[dict(name='接入点',values=[12,8,3])])
    add_chart(s,data,dict(x=Inches(.5),y=Inches(.5),width=Inches(8),height=Inches(5)))
    path=tmp_path/'chart.pptx';p.save(path);c=Presentation(path).slides[0].shapes[0].chart
    assert list(c.series[0].values)==[12,8,3] and c.series[0].name=='接入点'
    assert '2026年第三季度' in c.chart_title.text_frame.text and '个' in c.chart_title.text_frame.text
    with ZipFile(path) as z:assert any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in z.namelist())
