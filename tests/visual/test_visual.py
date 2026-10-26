import hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.validator.visual import visual_issues, contrast


def test_overlap_distortion_and_stale_render(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    for x in [1, 1.2]:
        a = s.shapes.add_textbox(Inches(x), Inches(1), Inches(4), Inches(1))
        a.text = "TEST FIXTURE overlap"
    source = tmp_path / "image.png"
    Image.new("RGB", (100, 50), "red").save(source)
    s.shapes.add_picture(str(source), Inches(1), Inches(3), Inches(3), Inches(3))
    file = tmp_path / "bad.pptx"
    p.save(file)
    png = tmp_path / "page.png"
    Image.new("RGB", (1000, 750), "white").save(png)
    rendered = {
        "sha256": hashlib.sha256(file.read_bytes()).hexdigest(),
        "page_count": 1,
        "images": [str(png)],
        "page_texts": ["TEST FIXTURE overlap"],
    }
    codes = {i.code for i in visual_issues(file, rendered)}
    assert {"TEXT_OVERLAP", "IMAGE_DISTORTION", "RENDER_BLANK"} <= codes
    rendered["sha256"] = "old"
    assert visual_issues(file, rendered)[0].code == "STALE_RENDER"


def test_contrast():
    assert contrast((255, 255, 255), (0, 0, 0)) > 20
    assert contrast((250, 250, 250), (255, 255, 255)) < 1.1
