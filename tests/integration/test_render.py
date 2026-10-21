from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches
from ppt_report_skill.renderer.render import render_presentation,find_office


@pytest.mark.render
def test_actual_office_render_and_cache(tmp_path):
    assert find_office(), 'RENDER_FAILED: 此验收要求真实 Office 渲染器'
    p=Presentation()
    for n in range(3):
        s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2));sh.text=f'TEST FIXTURE 第{n+1}页 完成率87.6%'
    source=tmp_path/'test.pptx';p.save(source)
    result=render_presentation(source,tmp_path/'rendered',dpi=80)
    assert result['page_count']==3 and len(result['images'])==3
    assert all(Path(i).stat().st_size>1000 for i in result['images'])
    assert all('87.6%' in t for t in result['page_texts'])
    timestamp=Path(result['pdf']).stat().st_mtime_ns
    assert render_presentation(source,tmp_path/'rendered',dpi=80)==result
    assert Path(result['pdf']).stat().st_mtime_ns==timestamp
