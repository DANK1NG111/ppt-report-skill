import importlib.util
from pathlib import Path
import pytest
from pptx import Presentation
from ppt_report_skill.validator.content import compact, slide_text


@pytest.mark.render
def test_second_template_full_pipeline(tmp_path):
    spec = importlib.util.spec_from_file_location(
        "second_template", Path("scripts/run_second_template.py")
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    result = module.run_case(tmp_path / "second_template")
    prs = Presentation(result["path"])
    text = compact("\n".join(slide_text(s) for s in prs.slides))
    assert all(
        value in text for value in ["87.6%", "12个监测点", "2026年9月30日", "2项未通过"]
    )
    assert all("Copyright TEST FIXTURE" in slide_text(slide) for slide in prs.slides)
    assert not any("正文示例材料" in slide_text(slide) for slide in prs.slides)
    assert result["rendered"]["page_count"] == 6
