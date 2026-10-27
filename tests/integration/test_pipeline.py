from pathlib import Path
import pytest
from pptx import Presentation
from ppt_report_skill.pipeline import create_presentation


@pytest.mark.render
def test_complete_pipeline_native_table(tmp_path):
    materials = [Path("tests/fixtures/materials/sample.xlsx")]
    result = create_presentation(
        "assets/templates/base_template.pptx",
        materials,
        {"topic": "TEST FIXTURE 原生数据汇报", "presentation_type": "custom"},
        tmp_path / "report.pptx",
    )
    assert result["status"] == "READY_FOR_VISUAL_REVIEW" and result["qa"]["errors"] == 0
    p = Presentation(result["path"])
    assert any(sh.has_table for s in p.slides for sh in s.shapes)
    assert result["rendered"]["page_count"] == len(p.slides)
    assert Path(result["build"], "source_facts.json").is_file()


def test_native_objects_stay_in_their_sections():
    from ppt_report_skill.pipeline import combine_plans
    from ppt_report_skill.models import PresentationRequest, SourceFact
    from ppt_report_skill.planner.outline import plan_outline
    from ppt_report_skill.template.analyzer import analyze_template

    profile = analyze_template("assets/templates/base_template.pptx")
    request = PresentationRequest(
        topic="TEST FIXTURE",
        presentation_type="custom",
        required_sections=["数据", "风险", "下一步"],
        visuals=[{"type": "chart", "section": "数据", "data": {"kind": "column"}}],
    )
    facts = [
        SourceFact(str(i), text, "fixture", str(i), "data", text, section=section)
        for i, (section, text) in enumerate(
            [
                ("数据", "本期接入12个监测点。"),
                ("风险", "存在数据延迟风险。"),
                ("下一步", "下期补充告警。"),
            ]
        )
    ]
    outline = plan_outline(request, facts, profile)
    plans = combine_plans(request, facts, outline, profile, [], 280, 6)
    positions = {
        section: [
            i
            for i, p in enumerate(plans)
            if p.section == section and p.slide_type == "content"
        ]
        for section in request.required_sections
    }
    assert max(positions["数据"]) < min(positions["风险"]) < min(positions["下一步"])


pytestmark = pytest.mark.private_template
