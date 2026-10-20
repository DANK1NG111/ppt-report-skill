import hashlib
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from ppt_report_skill.models import PresentationRequest
from ppt_report_skill.template.analyzer import analyze_template
from ppt_report_skill.content.loader import text_blocks
from ppt_report_skill.content.facts import extract_facts
from ppt_report_skill.planner.outline import plan_outline
from ppt_report_skill.planner.slide_planner import plan_slides
from ppt_report_skill.generator.presentation import generate_presentation


def test_real_template_generation(tmp_path):
    source='assets/templates/base_template.pptx';profile=analyze_template(source)
    request=PresentationRequest(topic='TEST FIXTURE 工作汇报',presentation_type='work_report',target_slide_count=8)
    facts=extract_facts(text_blocks('# 成果\n阶段完成率 87.6%，接入 12 个监测点。\n# 下一步\n计划开展第二轮测试。'))
    outline=plan_outline(request,facts,profile);plans=plan_slides(request,facts,outline,profile)
    out=generate_presentation(source,profile,plans,request,tmp_path/'report.pptx',tmp_path/'build')
    prs=Presentation(out)
    assert len(prs.slides)==8
    text='\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert all(f.content in text for f in facts)
    assert '单击此处添加' not in text and '87.6%' in text
    with ZipFile(source) as a,ZipFile(out) as b:
        for name in [n for n in a.namelist() if n.startswith('ppt/theme/') and n.endswith('.xml')]:assert a.read(name)==b.read(name)
        for name in [n for n in b.namelist() if n.startswith('ppt/slides/') and n.endswith('.xml')]:
            root=etree.fromstring(b.read(name))
            for paragraph in root.xpath('//*[local-name()="p"]'):
                assert len(paragraph.xpath('./*[local-name()="pPr"]'))<=1
    assert hashlib.sha256(open(source,'rb').read()).hexdigest()==profile.sha256
