from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from ppt_report_skill.repair import repair_loop
from ppt_report_skill.validator.content import slide_text


def test_real_render_repair_cycle(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    sh = s.shapes.add_textbox(Inches(-1), Inches(0.5), Inches(5), Inches(1))
    sh.text = "TEST FIXTURE 越界内容保留87.6%"
    sh = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(5), Inches(0.3))
    sh.text = "原始内容保持不变，2026-10-01完成第一轮测试，12个监测点全部接入。" * 2
    sh.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    s = p.slides.add_slide(p.slide_layouts[6])
    sh = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
    sh.text = "TEST FIXTURE 图片比例"
    image = tmp_path / "wide.png"
    Image.new("RGB", (800, 400), "#a6383b").save(image)
    s.shapes.add_picture(str(image), Inches(1), Inches(2), Inches(3), Inches(3))
    original = tmp_path / "bad.pptx"
    p.save(original)
    before = [slide_text(s) for s in p.slides]
    result = repair_loop(original, tmp_path / "repair")
    assert result["resolved"] and result["history"][0]["errors"] >= 3
    assert [slide_text(s) for s in Presentation(result["path"]).slides] == before
    assert result["rendered"]["page_count"] == 2
    assert (
        original.read_bytes() != __import__("pathlib").Path(result["path"]).read_bytes()
    )
